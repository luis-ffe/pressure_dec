"""Create a standalone interactive delay-analysis plot from a SensorTester workbook.

This is an experiment only; it is not imported by the SensorTester application.
"""

from __future__ import annotations

import argparse
import json
import statistics
from pathlib import Path

from openpyxl import load_workbook


def load_measurements(path: Path) -> tuple[list[float], list[float], list[float]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Measurements" not in workbook.sheetnames:
        raise ValueError("Workbook does not contain a 'Measurements' sheet")
    sheet = workbook["Measurements"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip().lower().replace(" ", "_") for value in next(rows)]
    def find_column(candidates: tuple[str, ...]) -> str:
        for candidate in candidates:
            if candidate in headers:
                return candidate
        raise ValueError(f"Missing required column; expected one of: {', '.join(candidates)}")

    time_column = find_column(("time_ms", "elapsed_ms", "elapsed_s", "time_s", "device_time_ms"))
    fsr1_column = find_column(("fsr1", "fsr1_raw", "f1"))
    fsr2_column = find_column(("fsr2", "fsr2_raw", "f2"))
    indexes = {
        "time": headers.index(time_column),
        "fsr1": headers.index(fsr1_column),
        "fsr2": headers.index(fsr2_column),
    }
    time_scale = 0.001 if time_column.endswith("_ms") else 1.0

    time_s: list[float] = []
    fsr1: list[float] = []
    fsr2: list[float] = []
    for row in rows:
        try:
            time_s.append(float(row[indexes["time"]]) * time_scale)
            fsr1.append(float(row[indexes["fsr1"]]))
            fsr2.append(float(row[indexes["fsr2"]]))
        except (TypeError, ValueError):
            continue
    if len(time_s) < 3:
        raise ValueError("At least three valid measurement rows are required")
    return time_s, fsr1, fsr2


def first_threshold_crossing(time_s: list[float], values: list[float], fraction: float = 0.10) -> tuple[float, float]:
    minimum = min(values)
    threshold = minimum + fraction * (max(values) - minimum)
    for timestamp, value in zip(time_s, values):
        if value >= threshold:
            return timestamp, threshold
    return time_s[-1], threshold


def create_html(source: Path, output: Path) -> dict[str, float]:
    time_s, fsr1, fsr2 = load_measurements(source)
    intervals = [later - earlier for earlier, later in zip(time_s, time_s[1:]) if later > earlier]
    median_interval = statistics.median(intervals)
    fsr1_onset, fsr1_threshold = first_threshold_crossing(time_s, fsr1)
    fsr2_onset, fsr2_threshold = first_threshold_crossing(time_s, fsr2)
    signed_delay = fsr2_onset - fsr1_onset
    delay_ms = signed_delay * 1000
    leader = "same sampled time"
    if signed_delay < 0:
        leader = "FSR 2 leads"
    elif signed_delay > 0:
        leader = "FSR 1 leads"

    payload = {
        "time": time_s,
        "fsr1": fsr1,
        "fsr2": fsr2,
        "fsr1Onset": fsr1_onset,
        "fsr2Onset": fsr2_onset,
        "fsr1Threshold": fsr1_threshold,
        "fsr2Threshold": fsr2_threshold,
        "delayMs": delay_ms,
        "medianIntervalMs": median_interval * 1000,
        "source": source.name,
    }
    data_json = json.dumps(payload, separators=(",", ":"))
    html = f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Force Sensor Delay Test</title>
  <script src="https://cdn.plot.ly/plotly-2.35.2.min.js"></script>
  <script src="https://cdn.jsdelivr.net/npm/xlsx@0.18.5/dist/xlsx.full.min.js"></script>
  <style>
    :root {{ --ink:#18212b; --muted:#607080; --grid:#dfe5ea; --blue:#2474b5; --orange:#d97723; }}
    * {{ box-sizing:border-box; }}
    body {{ margin:0; background:#f5f7f9; color:var(--ink); font:15px/1.45 -apple-system,BlinkMacSystemFont,"Segoe UI",sans-serif; }}
    main {{ max-width:1220px; margin:24px auto; padding:0 20px; }}
    h1 {{ margin:0 0 4px; font-size:28px; }}
    .subtitle {{ color:var(--muted); margin-bottom:16px; }}
    .toolbar {{ display:flex; align-items:center; gap:10px; flex-wrap:wrap; background:white; border:1px solid var(--grid); border-radius:10px; padding:12px 14px; margin-bottom:12px; }}
    .file-button {{ display:inline-flex; align-items:center; border:1px solid #1f659b; background:#2474b5; color:white; border-radius:7px; padding:8px 13px; cursor:pointer; font-weight:700; }}
    .file-button:hover {{ background:#1f659b; }}
    #file-input {{ position:absolute; width:1px; height:1px; overflow:hidden; clip:rect(0,0,0,0); }}
    #file-status {{ color:var(--muted); }}
    #file-status.error {{ color:#b42318; font-weight:650; }}
    .cards {{ display:grid; grid-template-columns:repeat(3,minmax(0,1fr)); gap:12px; margin-bottom:12px; }}
    .card {{ background:white; border:1px solid var(--grid); border-radius:10px; padding:12px 14px; }}
    .label {{ color:var(--muted); font-size:12px; font-weight:700; letter-spacing:.04em; text-transform:uppercase; }}
    .value {{ font-size:22px; font-weight:750; margin-top:2px; }}
    #chart {{ height:620px; background:white; border:1px solid var(--grid); border-radius:10px; }}
    .instructions {{ background:#fff8e9; border:1px solid #ead7ad; border-radius:10px; margin-top:12px; padding:10px 14px; }}
    button {{ border:1px solid #aebac4; background:white; border-radius:7px; padding:7px 11px; cursor:pointer; font-weight:600; }}
    @media(max-width:760px) {{ .cards {{ grid-template-columns:1fr; }} #chart {{ height:540px; }} }}
  </style>
</head>
<body>
<main>
  <h1>Force sensor delay test</h1>
  <div class="subtitle" id="subtitle">Interactive view of {source.name} · raw ESP32 ADC counts · {len(time_s)} recorded samples</div>
  <div class="toolbar">
    <label class="file-button" for="file-input">Open CSV or Excel file</label>
    <input id="file-input" type="file" accept=".csv,.xlsx,.xls,text/csv,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/vnd.ms-excel">
    <span id="file-status">Files are processed locally in this browser.</span>
  </div>
  <div class="cards">
    <div class="card"><div class="label">Automatic onset delay</div><div class="value" id="automatic-delay">{abs(delay_ms):.0f} ms</div><div id="automatic-detail">{leader} at the first 10% crossing</div></div>
    <div class="card"><div class="label">Sampling resolution</div><div class="value" id="sampling-resolution">{median_interval * 1000:.0f} ms</div><div>Delays smaller than roughly one sample are uncertain</div></div>
    <div class="card"><div class="label">Manual cursor delay</div><div class="value" id="manual-delay">Click twice</div><div id="cursor-readout">Choose two change points on either line</div></div>
  </div>
  <div id="chart"></div>
  <div class="instructions"><b>How to test:</b> hover for exact readings; drag to zoom; double-click to reset. Click one change point to set cursor A, then another to set cursor B. The measured Δt appears above. <button id="clear-cursors">Clear cursors</button></div>
</main>
<script>
let data = {data_json};
let cursors=[];
const chart=document.getElementById('chart');

function firstCrossing(time, values) {{
  let minimum=Infinity,maximum=-Infinity;
  for(const value of values) {{
    if(value<minimum) minimum=value;
    if(value>maximum) maximum=value;
  }}
  const threshold=minimum+0.10*(maximum-minimum);
  const index=values.findIndex(value=>value>=threshold);
  return {{time:time[index<0 ? time.length-1 : index],threshold}};
}}

function median(values) {{
  const ordered=[...values].sort((a,b)=>a-b);
  const middle=Math.floor(ordered.length/2);
  return ordered.length%2 ? ordered[middle] : (ordered[middle-1]+ordered[middle])/2;
}}

function analyze(time,fsr1,fsr2,source) {{
  const origin=time[0];
  const relativeTime=time.map(value=>value-origin);
  const intervals=relativeTime.slice(1).map((value,index)=>value-relativeTime[index]).filter(value=>value>0);
  const fsr1Crossing=firstCrossing(relativeTime,fsr1);
  const fsr2Crossing=firstCrossing(relativeTime,fsr2);
  return {{
    time:relativeTime,fsr1,fsr2,source,
    fsr1Onset:fsr1Crossing.time,fsr2Onset:fsr2Crossing.time,
    fsr1Threshold:fsr1Crossing.threshold,fsr2Threshold:fsr2Crossing.threshold,
    delayMs:(fsr2Crossing.time-fsr1Crossing.time)*1000,
    medianIntervalMs:intervals.length ? median(intervals)*1000 : 0
  }};
}}

function tracesFor(current) {{ return [
  {{x:current.time,y:current.fsr1,name:'FSR 1',type:'scattergl',mode:'lines',line:{{color:'#2474b5',width:2.2}},hovertemplate:'FSR 1: %{{y:.0f}}<br>Time: %{{x:.3f}} s<extra></extra>'}},
  {{x:current.time,y:current.fsr2,name:'FSR 2',type:'scattergl',mode:'lines',line:{{color:'#d97723',width:2.2,dash:'dash'}},hovertemplate:'FSR 2: %{{y:.0f}}<br>Time: %{{x:.3f}} s<extra></extra>'}},
  {{x:[current.fsr1Onset],y:[current.fsr1Threshold],name:'FSR 1 onset',type:'scatter',mode:'markers',marker:{{color:'#2474b5',size:11,symbol:'circle-open',line:{{width:2}}}},hovertemplate:'FSR 1 10% onset<br>%{{x:.3f}} s<extra></extra>'}},
  {{x:[current.fsr2Onset],y:[current.fsr2Threshold],name:'FSR 2 onset',type:'scatter',mode:'markers',marker:{{color:'#d97723',size:11,symbol:'diamond-open',line:{{width:2}}}},hovertemplate:'FSR 2 10% onset<br>%{{x:.3f}} s<extra></extra>'}}
]}};

function layoutFor(current) {{ return {{
  title:{{text:'FSR response over time<br><sup>Automatic markers use the first crossing of 10% of each channel’s recorded range</sup>',x:.04,y:.98,xanchor:'left',yanchor:'top'}},
  margin:{{l:72,r:28,t:118,b:90}},paper_bgcolor:'#fff',plot_bgcolor:'#fff',
  xaxis:{{title:'Elapsed time (seconds)',range:[current.time[0],current.time[current.time.length-1]],autorange:false,showgrid:true,gridcolor:'#e7ebef',rangeslider:{{visible:true,thickness:.08}},spikemode:'across',spikesnap:'cursor',showspikes:true,spikecolor:'#66727d'}},
  yaxis:{{title:'Raw ADC reading (0–4095)',range:[0,4300],showgrid:true,gridcolor:'#e7ebef',zeroline:false}},
  legend:{{orientation:'h',x:.02,y:1.015,xanchor:'left',yanchor:'bottom'}},hovermode:'x unified',dragmode:'zoom',
  shapes:[],annotations:[]
}}}};

function renderData(nextData) {{
  data=nextData;
  cursors=[];
  const magnitude=Math.abs(data.delayMs);
  const leader=magnitude<0.5 ? 'Both sensors respond together' : (data.delayMs>0 ? 'FSR 1 leads' : 'FSR 2 leads');
  document.getElementById('subtitle').textContent='Interactive view of '+data.source+' · raw ESP32 ADC counts · '+data.time.length+' recorded samples';
  document.getElementById('automatic-delay').textContent=magnitude.toFixed(0)+' ms';
  document.getElementById('automatic-detail').textContent=leader+' at the first 10% crossing';
  document.getElementById('sampling-resolution').textContent=data.medianIntervalMs.toFixed(1).replace('.0','')+' ms';
  Plotly.react(chart,tracesFor(data),layoutFor(data),{{responsive:true,displaylogo:false,scrollZoom:true}});
  renderCursors();
}}

function renderCursors() {{
  const colors=['#6b4c9a','#2f7d52'];
  const shapes=cursors.map((x,i)=>({{type:'line',xref:'x',yref:'paper',x0:x,x1:x,y0:0,y1:1,line:{{color:colors[i],width:2,dash:'dot'}}}}));
  Plotly.relayout(chart,{{shapes}});
  if(cursors.length===2) {{
    const delta=Math.abs(cursors[1]-cursors[0]);
    document.getElementById('manual-delay').textContent=(delta*1000).toFixed(0)+' ms';
    document.getElementById('cursor-readout').textContent='A = '+cursors[0].toFixed(3)+' s · B = '+cursors[1].toFixed(3)+' s · Δt = '+delta.toFixed(3)+' s';
  }} else if(cursors.length===1) {{
    document.getElementById('manual-delay').textContent='Cursor A set';
    document.getElementById('cursor-readout').textContent='A = '+cursors[0].toFixed(3)+' s · click cursor B';
  }} else {{
    document.getElementById('manual-delay').textContent='Click twice';
    document.getElementById('cursor-readout').textContent='Choose two change points on either line';
  }}
}}
renderData(data);

chart.on('plotly_click',event=>{{
  const x=Number(event.points[0].x);
  if(cursors.length===2) cursors=[];
  cursors.push(x); renderCursors();
}});
document.getElementById('clear-cursors').addEventListener('click',()=>{{cursors=[];renderCursors();}});

function normalizedHeader(value) {{
  return String(value??'').trim().toLowerCase().replace(/[^a-z0-9]+/g,'_').replace(/^_|_$/g,'');
}}

function findColumn(headers,candidates) {{
  for(const candidate of candidates) {{
    const index=headers.indexOf(candidate);
    if(index>=0) return index;
  }}
  return -1;
}}

function measurementsFromSheet(sheet) {{
  const matrix=XLSX.utils.sheet_to_json(sheet,{{header:1,raw:true,defval:null}});
  if(matrix.length<4) throw new Error('At least three measurement rows are required.');
  const headers=matrix.shift().map(normalizedHeader);
  const timeMsIndex=findColumn(headers,['time_ms','elapsed_ms','device_time_ms']);
  const timeSecondsIndex=findColumn(headers,['elapsed_s','time_s','elapsed_time_s','time_seconds']);
  const fsr1Index=findColumn(headers,['fsr1','fsr1_raw','f1']);
  const fsr2Index=findColumn(headers,['fsr2','fsr2_raw','f2']);
  const timeIndex=timeMsIndex>=0 ? timeMsIndex : timeSecondsIndex;
  if(timeIndex<0 || fsr1Index<0 || fsr2Index<0) {{
    throw new Error('Required columns: time_ms (or elapsed_s), fsr1, and fsr2.');
  }}
  const scale=timeMsIndex>=0 ? 0.001 : 1;
  const time=[],fsr1=[],fsr2=[];
  for(const row of matrix) {{
    const t=Number(row[timeIndex])*scale;
    const first=Number(row[fsr1Index]);
    const second=Number(row[fsr2Index]);
    if(Number.isFinite(t) && Number.isFinite(first) && Number.isFinite(second)) {{
      time.push(t);fsr1.push(first);fsr2.push(second);
    }}
  }}
  if(time.length<3) throw new Error('At least three valid numeric measurement rows are required.');
  return {{time,fsr1,fsr2}};
}}

document.getElementById('file-input').addEventListener('change',async event=>{{
  const file=event.target.files[0];
  if(!file) return;
  const status=document.getElementById('file-status');
  status.className='';
  status.textContent='Opening '+file.name+'…';
  try {{
    const workbook=XLSX.read(await file.arrayBuffer(),{{type:'array'}});
    const sheetName=workbook.SheetNames.includes('Measurements') ? 'Measurements' : workbook.SheetNames[0];
    const measurements=measurementsFromSheet(workbook.Sheets[sheetName]);
    renderData(analyze(measurements.time,measurements.fsr1,measurements.fsr2,file.name));
    status.textContent='Loaded '+file.name+' · '+measurements.time.length+' samples';
  }} catch(error) {{
    status.className='error';
    status.textContent=error instanceof Error ? error.message : String(error);
  }} finally {{
    event.target.value='';
  }}
}});
</script>
</body>
</html>"""
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(html, encoding="utf-8")
    return {
        "samples": float(len(time_s)),
        "median_interval_ms": median_interval * 1000,
        "fsr1_onset_s": fsr1_onset,
        "fsr2_onset_s": fsr2_onset,
        "signed_delay_ms": delay_ms,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, default=Path("force_delay_analysis.html"))
    args = parser.parse_args()
    result = create_html(args.workbook, args.output)
    print(f"Created {args.output.resolve()}")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
