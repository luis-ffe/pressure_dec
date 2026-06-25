import csv
import struct
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

from sensor_tester.communications import Sample, SerialTransport, WifiTransport
from sensor_tester.data_export import (
    FIELDNAMES,
    Measurement,
    calculate_response_delay,
    export_csv,
    export_xlsx,
)


class SampleTests(unittest.TestCase):
    def test_accepts_usb_payload(self):
        sample = Sample.from_payload(
            {"time_ms": 123, "fsr1": 100, "fsr2": 200, "moving": True, "steps_remaining": 42}
        )
        self.assertEqual(sample.fsr1_raw, 100)
        self.assertEqual(sample.steps_remaining, 42)

    def test_accepts_legacy_wifi_payload(self):
        sample = Sample.from_payload({"time": 123, "f1": 11, "f2": 22})
        self.assertEqual((sample.fsr1_raw, sample.fsr2_raw), (11, 22))
        self.assertFalse(sample.motion_status_known)

    def test_accepts_compact_high_rate_usb_sample(self):
        sample = Sample.from_compact_usb("S,1234,111,222")
        self.assertEqual(sample.device_time_ms, 1234)
        self.assertEqual(sample.device_time_us, 1_234_000)
        self.assertEqual((sample.fsr1_raw, sample.fsr2_raw), (111, 222))
        self.assertFalse(sample.motion_status_known)


class TransportTests(unittest.TestCase):
    def test_serial_defaults_to_dma_baudrate(self):
        self.assertEqual(SerialTransport("loop://").baudrate, 460_800)

    def test_serial_decodes_fixed_binary_dma_capture(self):
        transport = SerialTransport("loop://")
        fake_serial = MagicMock()
        pairs = [(index & 0x0FFF, (index + 100) & 0x0FFF) for index in range(5_000)]
        payload = b"".join(struct.pack("<HH", *pair) for pair in pairs)
        fake_serial.in_waiting = len(payload)
        fake_serial.read.return_value = payload
        transport._serial = fake_serial

        first = transport.read_sample()
        second = transport.read_sample()

        self.assertIsNotNone(first)
        self.assertIsNotNone(second)
        assert first is not None and second is not None
        self.assertEqual((first.fsr1_raw, first.fsr2_raw), pairs[0])
        self.assertEqual(first.sequence, 0)
        self.assertEqual(second.device_time_us - first.device_time_us, 100)
        self.assertEqual(len(transport._decoded_samples), 4_998)

    def test_serial_move_command(self):
        transport = SerialTransport("loop://")
        fake_serial = MagicMock()
        transport._serial = fake_serial
        transport.move("up", 200, 500)
        fake_serial.write.assert_called_once_with(b"MOVE,UP,200,500\n")

    def test_wifi_maps_down_to_backward(self):
        transport = WifiTransport("http://192.168.4.1")
        with patch.object(transport, "_get") as request:
            transport.move("down", 100, 700)
        request.assert_called_once_with("/move", dir="backward", steps=100, speed=700)

    def test_serial_recording_mode_commands(self):
        transport = SerialTransport("loop://")
        fake_serial = MagicMock()
        transport._serial = fake_serial
        transport.set_recording(True)
        transport.set_recording(False)
        self.assertEqual(
            [call.args[0] for call in fake_serial.write.call_args_list],
            [b"RECORD,1\n", b"RECORD,0\n"],
        )


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.rows = [
            Measurement(
                time_ms=200,
                fsr1=111,
                fsr2=222,
            )
        ]

    def test_csv_export(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "data.csv"
            export_csv(path, self.rows)
            with path.open(encoding="utf-8-sig", newline="") as handle:
                exported = list(csv.DictReader(handle))
            self.assertEqual(list(exported[0]), FIELDNAMES)
            self.assertEqual(FIELDNAMES, ["time_ms", "fsr1", "fsr2"])
            self.assertEqual(exported[0]["fsr2"], "222")

    def test_xlsx_export(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "data.xlsx"
            export_xlsx(path, self.rows)
            workbook = load_workbook(path, data_only=False)
            self.assertEqual(workbook.sheetnames, ["Measurements", "Summary"])
            self.assertEqual(workbook["Measurements"]["A2"].value, 200)
            self.assertEqual(workbook["Measurements"]["B2"].value, 111)
            self.assertEqual(len(workbook["Summary"]._charts), 1)

    def test_response_delay_uses_first_ten_percent_crossing(self):
        rows = []
        for index, (fsr1, fsr2) in enumerate([(0, 0), (150, 0), (500, 200), (900, 800)]):
            rows.append(
                Measurement(
                    time_ms=index * 4,
                    fsr1=fsr1,
                    fsr2=fsr2,
                )
            )
        delay = calculate_response_delay(rows)
        self.assertIsNotNone(delay)
        assert delay is not None
        self.assertEqual(delay.signed_delay_ms, 4.0)
        self.assertIn("FSR 1 leads", delay.summary)


if __name__ == "__main__":
    unittest.main()
